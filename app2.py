import streamlit as st
import fitz  # PyMuPDF
import docx
import ollama
import tempfile
import speech_recognition as sr
import os
import re
import pandas as pd
import json
from datetime import datetime
from fpdf import FPDF
from streamlit_mic_recorder import mic_recorder
import soundfile as sf
import io

# ==============================================================
# UTILS: PEMBERSIH & PARSING
# ==============================================================
def sanitize_for_pdf(text):
    if not isinstance(text, str): text = str(text)
    text = text.replace("μ", "u").replace("’", "'").replace("”", '"').replace("“", '"').replace("…", "...")
    text = text.replace("\n", " ").replace("\r", "")
    return text.encode('latin-1', 'replace').decode('latin-1')

def clean_filename(filename):
    return re.sub(r'[\\/*?:"<>|]', "", filename).strip()

def extract_json_from_text(text):
    """Ekstrak JSON objek tunggal (untuk info pasien)"""
    try:
        start_obj = text.find('{')
        end_obj = text.rfind('}') + 1
        if start_obj != -1 and end_obj != -1:
             try: return json.loads(text[start_obj:end_obj])
             except: pass
        return {}
    except: return {}

# ==============================================================
# FUNGSI EKSTRAKSI DOKUMEN
# ==============================================================
def extract_text_from_pdf(file_path, password=None):
    try:
        doc = fitz.open(file_path)
        if doc.is_encrypted:
            if not password: return "ERROR: Password required"
            if not doc.authenticate(password): return "ERROR: Invalid password"
        
        text = ""
        for page in doc:
            text += page.get_text("text", sort=True) + "\n"
        return text.strip()
    except Exception as e:
        if "encrypted" in str(e).lower(): return "ERROR: Password required"
        return f"ERROR: {str(e)}"

def extract_text_from_docx(file_path):
    doc = docx.Document(file_path)
    text = "\n".join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells])
            text += "\n" + row_text
    return text.strip()

# ==============================================================
# FUNGSI AUDIO
# ==============================================================
def transcribe_audio_bytes(audio_bytes):
    recognizer = sr.Recognizer()
    temp_filename = f"rec_{datetime.now().strftime('%H%M%S%f')}.wav"
    temp_path = os.path.join(tempfile.gettempdir(), temp_filename)
    
    try:
        data, samplerate = sf.read(io.BytesIO(audio_bytes))
        sf.write(temp_path, data, samplerate, subtype='PCM_16')
        
        with sr.AudioFile(temp_path) as source:
            audio_data = recognizer.record(source)
            text = recognizer.recognize_google(audio_data, language="id-ID")
            return text
    except sr.UnknownValueError:
        return ""
    except sr.RequestError:
        return "Error: Koneksi internet bermasalah (Google API)."
    except Exception as e:
        return f"Error: {str(e)}"
    finally:
        if os.path.exists(temp_path):
            try: os.remove(temp_path)
            except: pass

# ==============================================================
# FUNGSI AI LOGIC
# ==============================================================
def extract_dashboard_info(text, model_name="mistral"):
    prompt = f"""
    Ekstrak data ADMINISTRATIF pasien berikut ke JSON Object.
    Keys: "Nama Pasien", "Tanggal Lahir", "Usia", "Gejala Utama", "No. Lab", "Tanggal Pemeriksaan", "Klinik/RS".
    Isi "Tidak ditemukan" jika data tidak ada.
    Output HANYA JSON.
    
    Teks: {text[:3000]} 
    """
    try:
        response = ollama.chat(model=model_name, messages=[{"role": "user", "content": prompt}], options={'temperature': 0.1})
        data = extract_json_from_text(response["message"]["content"])
        if isinstance(data, dict): return data
        return {} 
    except: return {}

def evaluate_status(val, ref):
    if not val or not ref: return "-"
    val = str(val).lower().replace(',', '.')
    ref = str(ref).lower().replace(',', '.')
    
    if "negatif" in ref:
        if "negatif" in val: return "Normal"
        if "positif" in val: return "Tinggi"
    
    try:
        v_match = re.search(r"(\d+\.?\d*)", val)
        if not v_match: return "-"
        v = float(v_match.group(1))
        
        if "-" in ref:
            nums = re.findall(r"(\d+\.?\d*)", ref)
            if len(nums) >= 2:
                low, high = float(nums[0]), float(nums[1])
                return "Normal" if low <= v <= high else ("Rendah" if v < low else "Tinggi")
        if "<" in ref:
            lim = float(re.search(r"(\d+\.?\d*)", ref).group(1))
            return "Normal" if v < lim else "Tinggi"
        if ">" in ref:
            lim = float(re.search(r"(\d+\.?\d*)", ref).group(1))
            return "Normal" if v > lim else "Rendah"
    except: pass
    return "-"

def extract_lab_data_with_ref(text, model_name="mistral"):
    prompt = f"""
    Tugas: Ekstrak SEMUA baris hasil pemeriksaan laboratorium dari teks.
    
    Format Output WAJIB setiap baris:
    NAMA PEMERIKSAAN | HASIL | SATUAN | NILAI RUJUKAN
    
    Contoh Output yang diinginkan:
    Hemoglobin | 13.5 | g/dL | 13.0 - 17.0
    
    Aturan:
    1. Abaikan baris judul.
    2. Jika satuan tidak ada, kosongkan.
    3. Gunakan pemisah tanda kurung tegak (|).
    4. Tuliskan SEMUA data yang ada nilainya.
    
    Teks Laporan:
    {text}
    """
    try:
        response = ollama.chat(model=model_name, messages=[{"role": "user", "content": prompt}], options={'temperature': 0.1})
        raw_text = response["message"]["content"]
        
        parsed_data = []
        for line in raw_text.split('\n'):
            line = line.strip()
            if not line or "|" not in line: continue
            if "NAMA PEMERIKSAAN" in line.upper(): continue 
            
            parts = [p.strip() for p in line.split('|')]
            
            if len(parts) >= 2:
                item = {
                    "Lab": parts[0],
                    "Nilai": parts[1],
                    "Unit": parts[2] if len(parts) > 2 else "",
                    "Nilai Rujukan": parts[3] if len(parts) > 3 else ""
                }
                item['Status'] = evaluate_status(item['Nilai'], item['Nilai Rujukan'])
                parsed_data.append(item)
                
        return parsed_data
    except: return []

def summarize_lab_results(lab_data, model_name="mistral"):
    if not lab_data: return "Tidak ada data lab."
    try:
        abnormal = [x for x in lab_data if x.get('Status') in ['Tinggi', 'Rendah'] or "positif" in str(x.get("Nilai", "")).lower()]
        if not abnormal: return "Berdasarkan hasil analisis, semua parameter dalam batas normal."
        data_str = ", ".join([f"{x['Lab']} ({x['Nilai']} {x['Unit']})" for x in abnormal])
    except: return "Error pemrosesan data."
    
    prompt = f"""
    Bertindaklah sebagai dokter ahli. Berikan ringkasan naratif singkat (maksimal 3 kalimat) dalam Bahasa Indonesia.
    Data abnormal pasien: {data_str}
    Jelaskan potensi indikasi medis dari data tersebut.
    """
    try:
        return ollama.chat(model=model_name, messages=[{"role": "user", "content": prompt}], options={'temperature': 0.4})["message"]["content"].strip()
    except: return "Gagal membuat ringkasan."

def answer_question(pkg, q, model_name="mistral"):
    try:
        ctx = f"Pasien: {pkg['name']}\nRingkasan: {pkg['lab_summary']}\n"
        ctx += "Data Lab:\n" + "\n".join([f"- {d['Lab']}: {d['Nilai']} (Ref: {d['Nilai Rujukan']})" for d in pkg['lab_data_list']])
        prompt = f"Jawab pertanyaan ini berdasarkan data medis berikut dalam Bahasa Indonesia:\n\nData:\n{ctx}\n\nPertanyaan: {q}"
        return ollama.chat(model=model_name, messages=[{"role": "user", "content": prompt}], options={'temperature': 0.5})["message"]["content"].strip()
    except: return "Error menjawab."

# ==============================================================
# PDF GENERATOR
# ==============================================================
class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 14)
        self.cell(0, 10, 'LAPORAN ANALISIS MEDIS', 0, 1, 'C')
        self.line(10, 20, 200, 20)
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'Halaman {self.page_no()}', 0, 0, 'R')
    
    def create_table(self, data):
        self.set_font('Arial', 'B', 9)
        w = [65, 30, 20, 25, 50]
        headers = ['Pemeriksaan', 'Hasil', 'Unit', 'Status', 'Rujukan']
        self.set_fill_color(240, 240, 240)
        for i, h in enumerate(headers):
            self.cell(w[i], 8, h, 1, 0, 'C', True)
        self.ln()
        self.set_font('Arial', '', 9)
        for row in data:
            vals = [
                str(row.get('Lab', '-')),
                str(row.get('Nilai', '-')),
                str(row.get('Unit', '')),
                str(row.get('Status', '-')),
                str(row.get('Nilai Rujukan', '-'))
            ]
            for i, txt in enumerate(vals):
                clean_txt = sanitize_for_pdf(txt)
                limit = {0: 45, 1: 18, 2: 10, 3: 15, 4: 30}.get(i, 20)
                if len(clean_txt) > limit: clean_txt = clean_txt[:limit] + ".."
                self.cell(w[i], 7, clean_txt, 1, 0, 'L')
            self.ln()

def generate_pdf_report(pkg):
    pdf = PDF('P', 'mm', 'A4')
    pdf.set_margins(10, 10, 10)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    
    pdf.set_font('Arial', 'B', 11)
    pdf.cell(0, 8, '1. INFORMASI PASIEN', 0, 1)
    pdf.set_font('Arial', '', 10)
    info_keys = [("Nama Pasien", "Nama Pasien"), ("Usia", "Usia"), ("Tanggal Lahir", "Tanggal Lahir"), ("Klinik/RS", "Klinik/RS")]
    info = pkg.get("patient_info_full", {})
    if not isinstance(info, dict): info = {}
    for label, key in info_keys:
        val = str(info.get(key, "-"))
        pdf.cell(45, 7, label, 0, 0)
        pdf.cell(5, 7, ":", 0, 0)
        pdf.cell(0, 7, sanitize_for_pdf(val), 0, 1)
    pdf.ln(5)
    
    pdf.set_font('Arial', 'B', 11)
    pdf.cell(0, 8, '2. RINGKASAN KLINIS', 0, 1)
    pdf.set_font('Arial', '', 10)
    summary_txt = sanitize_for_pdf(pkg.get("lab_summary", "-"))
    pdf.multi_cell(0, 6, summary_txt)
    pdf.ln(8)
    
    pdf.set_font('Arial', 'B', 11)
    pdf.cell(0, 8, '3. DETAIL HASIL LABORATORIUM', 0, 1)
    if pkg.get("lab_data_list") and len(pkg["lab_data_list"]) > 0:
        pdf.create_table(pkg["lab_data_list"])
    else:
        pdf.set_font('Arial', 'I', 10)
        pdf.cell(0, 8, "Tidak ada data laboratorium detail.", 0, 1)
    
    return bytes(pdf.output(dest='S'))

# ==============================================================
# MAIN UI
# ==============================================================
st.set_page_config(page_title="Medical-AI-Summarizer", layout="wide")

keys = ['unlocked_text','current_file_name','history','current_idx','chat_hist','last_uploaded_filename', 'last_audio_bytes']
for k in keys:
    if k not in st.session_state: st.session_state[k] = "" if "name" in k or "text" in k else [] if "hist" in k else None
if 'password_needed' not in st.session_state: st.session_state.password_needed = False
if 'temp_path' not in st.session_state: st.session_state.temp_path = ""

st.title("Medical-AI-Summarizer")

with st.sidebar:
    st.header("Input Data")
    method = st.radio("Metode:", ["Unggah Dokumen", "Manual", "Suara"])
    
    if method == "Manual":
        if t := st.text_area("Paste Teks:", height=200):
            st.session_state.unlocked_text = t
            st.session_state.password_needed = False

    elif method == "Suara":
        st.info("Klik tombol di bawah, bicara, lalu hentikan.")
        aud = mic_recorder(start_prompt="Mulai Rekam", stop_prompt="Stop", key='mic_recorder', format="wav")
        
        if aud and aud.get('bytes'):
            if aud['bytes'] != st.session_state.last_audio_bytes:
                st.session_state.last_audio_bytes = aud['bytes']
                
                with st.spinner("Menerjemahkan suara ke teks..."):
                    text = transcribe_audio_bytes(aud['bytes'])
                    if text:
                        st.session_state.unlocked_text = text
                        st.session_state.password_needed = False
                        st.success("Berhasil!")
                        st.rerun()
                    else:
                        st.warning("Suara tidak terdengar, coba lagi.")

    elif method == "Unggah Dokumen":
        f = st.file_uploader("PDF/DOCX:", type=["pdf", "docx"])
        if f:
            if f.name != st.session_state.last_uploaded_filename:
                st.session_state.last_uploaded_filename = f.name
                st.session_state.current_file_name = f.name
                st.session_state.unlocked_text = ""
                st.session_state.password_needed = False
                
                path = os.path.join(tempfile.gettempdir(), f.name)
                with open(path, "wb") as wb: wb.write(f.read())
                st.session_state.temp_path = path
                
                if f.name.endswith(".pdf"):
                    res = extract_text_from_pdf(path)
                    if "Password required" in res:
                        st.session_state.password_needed = True
                        st.warning("Terkunci Password")
                    elif "ERROR" in res: st.error(res)
                    else: st.session_state.unlocked_text = res
                else:
                    st.session_state.unlocked_text = extract_text_from_docx(path)

        if st.session_state.password_needed:
            pwd = st.text_input("Password:", type="password")
            if pwd:
                with st.spinner("Membuka..."):
                    res = extract_text_from_pdf(st.session_state.temp_path, pwd)
                    if "Invalid" in res: st.error("Salah Password")
                    elif "ERROR" in res: st.error(res)
                    else:
                        st.session_state.unlocked_text = res
                        st.session_state.password_needed = False
                        st.rerun()

    if st.session_state.unlocked_text and not st.session_state.password_needed:
        st.divider()
        if st.button("Analisis Dokumen", type="primary", use_container_width=True):
            with st.spinner("AI sedang membaca data..."):
                txt = st.session_state.unlocked_text
                info = extract_dashboard_info(txt)
                if not isinstance(info, dict): info = {}
                labs = extract_lab_data_with_ref(txt)
                summ = summarize_lab_results(labs)
                
                name = info.get("Nama Pasien", "Pasien")
                pkg = {
                    "name": f"{name}",
                    "patient_info_full": info,
                    "lab_data_list": labs,
                    "lab_summary": summ,
                    "full_text": txt
                }
                st.session_state.history.append(pkg)
                st.session_state.current_idx = len(st.session_state.history) - 1
                st.rerun()

    # HISTORY & MAIN CONTENT
    if st.session_state.history:
        st.divider()
        st.subheader("Arsip")
        opts = [h['name'] for h in st.session_state.history]
        idx = st.selectbox("Pilih:", range(len(opts)), format_func=lambda x: opts[x], index=st.session_state.current_idx)
        
        if idx != st.session_state.current_idx:
            st.session_state.current_idx = idx
            st.rerun()
            
        curr = st.session_state.history[idx]
        try:
            curr_info = curr.get('patient_info_full', {})
            if not isinstance(curr_info, dict): curr_info = {}
            p_name = clean_filename(curr_info.get("Nama Pasien", "Pasien"))
            fname = f"{p_name}_Hasil_Lab.pdf"
            pdf_bytes = generate_pdf_report(curr)
            st.download_button("Unduh PDF", pdf_bytes, fname, "application/pdf", use_container_width=True)
        except Exception as e: st.error(f"PDF Error: {e}")
        
        if st.button("Hapus"):
            del st.session_state.history[idx]
            st.session_state.current_idx = 0 if st.session_state.history else None
            st.rerun()

if st.session_state.history and st.session_state.current_idx is not None:
    curr = st.session_state.history[st.session_state.current_idx]
    inf = curr.get("patient_info_full", {})
    if not isinstance(inf, dict): inf = {}
    
    st.subheader(f"Analisis: {inf.get('Nama Pasien', 'Pasien')}")
    
    c1, c2, c3 = st.columns(3)
    c1.info(f"**Usia:** {inf.get('Usia', '-')}")
    c2.info(f"**Tgl:** {inf.get('Tanggal Pemeriksaan', '-')}")
    c3.info(f"**RS:** {inf.get('Klinik/RS', '-')}")
    
    tab1, tab2, tab3 = st.tabs(["Ringkasan", "Tabel Data", "Tanya AI"])
    
    with tab1:
        st.markdown("### Kesimpulan Medis")
        st.success(curr.get("lab_summary", "-"))
        st.divider()
        st.markdown("### Data Pasien")
        st.dataframe(pd.DataFrame(inf.items(), columns=["Parameter", "Nilai"]), hide_index=True, use_container_width=True)
        
    with tab2:
        labs = curr.get("lab_data_list", [])
        if labs:
            df = pd.DataFrame(labs)
            cols = ["Lab", "Nilai", "Unit", "Status", "Nilai Rujukan"]
            st.dataframe(df[[c for c in cols if c in df.columns]], hide_index=True, use_container_width=True)
        else:
            st.warning("Data detail laboratorium tidak terbaca.")
            
    with tab3:
        for m in st.session_state.chat_hist:
            st.chat_message(m["role"]).write(m["content"])
        if q := st.chat_input("Tanya tentang hasil ini..."):
            st.session_state.chat_hist.append({"role": "user", "content": q})
            st.chat_message("user").write(q)
            with st.spinner("..."):
                ans = answer_question(curr, q)
                st.session_state.chat_hist.append({"role": "assistant", "content": ans})
                st.chat_message("assistant").write(ans)
else:
    st.info("Silakan pilih metode input di sidebar.")